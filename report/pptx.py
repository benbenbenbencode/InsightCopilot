"""PowerPoint deck rendering with python-pptx.

Six slides: title, executive summary + KPIs, chart, findings, actions, and a
data/provenance appendix. Built from the same ReportBundle as the PDF.
"""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from core.analysis import fmt_number
from report.model import ReportBundle

INK = RGBColor(0x18, 0x18, 0x20)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x8A, 0x8A, 0x99)
ACCENT = RGBColor(0x6C, 0x5C, 0xE7)
ACCENT_SOFT = RGBColor(0xEE, 0xEB, 0xFF)
PANEL = RGBColor(0xF6, 0xF6, 0xFA)
BORDER = RGBColor(0xE2, 0xE2, 0xE8)
POSITIVE = RGBColor(0x00, 0x96, 0x78)
WARNING = RGBColor(0xC4, 0x78, 0x14)
CRITICAL = RGBColor(0xBE, 0x32, 0x32)

SEVERITY = {"positive": POSITIVE, "warning": WARNING, "critical": CRITICAL, "info": ACCENT}

W, H = 13.333, 7.5


def _rect(slide, x, y, w, h, color, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    if rounded:
        try:
            shape.adjustments[0] = 0.06
        except Exception:
            pass
    return shape


def _text(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
          font="Segoe UI", spacing=1.0, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0

    paragraphs = str(text).split("\n")
    for index, line in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.line_spacing = spacing
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = font
    return box


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide, title, kicker="RETAIL INSIGHT"):
    _text(slide, 0.6, 0.42, 6, 0.3, kicker, size=9, bold=True, color=ACCENT)
    _text(slide, 0.6, 0.7, 12, 0.6, title, size=26, bold=True, color=INK)


def _footer(slide, bundle: ReportBundle, page: int):
    _text(slide, 0.6, 7.02, 9, 0.3, f"Generated {bundle.stamp} with a local model - {bundle.narrative.model}",
          size=8, color=MUTED)
    _text(slide, 11.9, 7.02, 0.9, 0.3, str(page), size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def _kpi_row(slide, kpis, y=2.6):
    cards = kpis[:4]
    if not cards:
        return
    gap = 0.3
    width = (W - 1.2 - gap * (len(cards) - 1)) / len(cards)
    x = 0.6

    for kpi in cards:
        _rect(slide, x, y, width, 1.35, PANEL)
        _text(slide, x + 0.22, y + 0.18, width - 0.44, 0.3, kpi.label.upper(), size=8.5, bold=True, color=MUTED)
        _text(slide, x + 0.22, y + 0.55, width - 0.44, 0.6, kpi.display, size=24, bold=True, color=INK)
        x += width + gap


def build_pptx(bundle: ReportBundle, path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # ---- 1. title --------------------------------------------------------
    slide = _blank(prs)
    _rect(slide, 0, 0, W, H, INK, rounded=False)
    _rect(slide, 0.6, 2.1, 0.09, 1.9, ACCENT, rounded=False)

    _text(slide, 1.0, 2.05, 11, 0.5, "RETAIL INSIGHT COPILOT", size=11, bold=True, color=ACCENT)
    _text(slide, 1.0, 2.55, 11.2, 1.2, bundle.narrative.headline or bundle.title, size=32, bold=True, color=WHITE, spacing=0.95)
    _text(slide, 1.0, 3.95, 11.2, 0.9, f"Q: {bundle.question}", size=13, color=RGBColor(0xB8, 0xB8, 0xC4))
    _text(
        slide, 1.0, 6.55, 11.2, 0.4,
        f"{bundle.stamp}   |   {bundle.analysis.rows} rows   |   local model: {bundle.narrative.model}",
        size=10, color=RGBColor(0x80, 0x80, 0x90),
    )

    # ---- 2. executive summary -------------------------------------------
    slide = _blank(prs)
    _header(slide, "Executive summary")
    _kpi_row(slide, bundle.analysis.kpis, y=1.65)

    _text(slide, 0.6, 3.35, 12.1, 1.1, bundle.narrative.summary, size=14, color=INK, spacing=1.15)

    points = bundle.narrative.key_points[:3]
    if points:
        y = 4.75
        for point in points:
            _rect(slide, 0.6, y + 0.09, 0.11, 0.11, ACCENT)
            _text(slide, 0.92, y, 11.6, 0.5, point, size=11.5, color=RGBColor(0x44, 0x44, 0x50))
            y += 0.62

    _footer(slide, bundle, 2)

    # ---- 3. chart --------------------------------------------------------
    if bundle.chart_png:
        slide = _blank(prs)
        _header(slide, bundle.chart_spec.get("title", "Visual"))

        # Fit the image inside the slide: preserve aspect ratio and cap the
        # height so the caption and footer never collide with the plot.
        head = bundle.chart_png[:24]
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            pixel_w = int.from_bytes(head[16:20], "big")
            pixel_h = int.from_bytes(head[20:24], "big")
            aspect = pixel_w / pixel_h if pixel_h else 2.0
        else:
            aspect = 2.0

        max_w, max_h = 11.6, 4.7
        width = min(max_w, max_h * aspect)
        height = width / aspect
        left = (W - width) / 2
        top = 1.55

        slide.shapes.add_picture(
            io.BytesIO(bundle.chart_png), Inches(left), Inches(top), width=Inches(width)
        )

        _text(
            slide,
            left,
            top + height + 0.14,
            width,
            0.3,
            f"{bundle.analysis.rows} rows  ·  {bundle.chart_spec.get('kind', '')} chart",
            size=9,
            color=MUTED,
        )
        _footer(slide, bundle, 3)

    # ---- 4. findings -----------------------------------------------------
    if bundle.analysis.findings:
        slide = _blank(prs)
        _header(slide, "What the data shows")
        y = 1.6
        for finding in bundle.analysis.findings[:5]:
            color = SEVERITY.get(finding.severity, ACCENT)
            _rect(slide, 0.6, y, 0.075, 0.82, color, rounded=False)
            _text(slide, 0.82, y, 11.5, 0.3, finding.title, size=13, bold=True, color=INK)
            _text(slide, 0.82, y + 0.32, 11.5, 0.5, finding.detail, size=10, color=MUTED)
            y += 1.02
        _footer(slide, bundle, 4)

    # ---- 5. actions ------------------------------------------------------
    slide = _blank(prs)
    _header(slide, "Recommended actions")

    y = 1.6
    for index, item in enumerate(bundle.narrative.recommendations[:4], 1):
        _rect(slide, 0.6, y, 0.46, 0.46, ACCENT_SOFT)
        _text(slide, 0.6, y + 0.08, 0.46, 0.3, str(index), size=12, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        _text(slide, 1.24, y + 0.04, 11.1, 0.6, item, size=11.5, color=INK)
        y += 0.78

    if bundle.narrative.risks:
        _rect(slide, 0.6, 5.15, 12.1, 1.55, RGBColor(0xFD, 0xF6, 0xEA))
        _text(slide, 0.85, 5.32, 11.6, 0.3, "CAVEATS", size=8.5, bold=True, color=WARNING)
        _text(slide, 0.85, 5.62, 11.6, 0.95, "\n".join(bundle.narrative.risks[:3]), size=9.5, color=RGBColor(0x70, 0x5A, 0x30))

    _footer(slide, bundle, 5)

    # ---- 6. data appendix ------------------------------------------------
    slide = _blank(prs)
    _header(slide, "Result data")

    df = bundle.df.head(9)
    columns = list(df.columns)[:6]
    if len(columns) and len(df):
        rows, cols = len(df) + 1, len(columns)
        table_shape = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.4 * rows))
        table = table_shape.table

        for index, column in enumerate(columns):
            cell = table.cell(0, index)
            cell.text = str(column)[:24]
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT_SOFT
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.bold = True
                    run.font.color.rgb = INK
                    run.font.name = "Segoe UI"

        for row_index, (_, row) in enumerate(df.iterrows(), start=1):
            for col_index, column in enumerate(columns):
                value = row[column]
                text = fmt_number(value) if isinstance(value, (int, float)) else str(value)
                cell = table.cell(row_index, col_index)
                cell.text = text[:30]
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x3D)
                        run.font.name = "Segoe UI"

        _text(slide, 0.6, 6.55, 11.8, 0.3,
              f"Showing {len(df)} of {bundle.analysis.rows} rows. SQL in the Excel export.",
              size=8.5, color=MUTED)
    else:
        _text(slide, 0.6, 2.0, 12, 0.5, "No rows returned for this question.", size=13, color=MUTED)

    _footer(slide, bundle, 6)

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output
