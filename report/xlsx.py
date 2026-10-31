"""Excel export: summary, KPIs, findings, raw data and the query itself."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from report.model import ReportBundle

HEADER_FILL = PatternFill("solid", fgColor="EEEAFF")
HEADER_FONT = Font(bold=True, color="241824", name="Segoe UI", size=10)
BODY_FONT = Font(color="33333D", name="Segoe UI", size=10)
TITLE_FONT = Font(bold=True, color="6C5CE7", name="Segoe UI", size=14)
THIN = Side(style="thin", color="E2E2E8")
BORDER = Border(bottom=THIN)


def _style_sheet(worksheet, width_cap: int = 42):
    """Header styling, column widths and frozen header row."""
    for cell in worksheet[1]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER

    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = 0
        for cell in column_cells[:200]:
            value = "" if cell.value is None else str(cell.value)
            longest = max(longest, len(value))
        worksheet.column_dimensions[letter].width = min(max(11, longest + 3), width_cap)

    worksheet.freeze_panes = "A2"


def build_xlsx(bundle: ReportBundle, path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)

    summary_rows = [
        ("Question", bundle.question),
        ("Headline", bundle.narrative.headline),
        ("Summary", bundle.narrative.summary),
        ("Rows returned", bundle.analysis.rows),
        ("Generated", bundle.stamp),
        ("Model", bundle.narrative.model),
        ("Grounded", "yes" if bundle.narrative.grounded else "no"),
        ("SQL", bundle.sql),
    ]
    summary = pd.DataFrame(summary_rows, columns=["Field", "Value"])

    kpis = pd.DataFrame(
        [{"Metric": k.label, "Value": k.value, "Display": k.display} for k in bundle.analysis.kpis]
    )

    findings = pd.DataFrame(
        [
            {
                "Severity": f.severity,
                "Finding": f.title,
                "Detail": f.detail,
                "Value": f.value,
            }
            for f in bundle.analysis.findings
        ]
    )

    # Recommendations and caveats have independent counts; build explicit rows
    # instead of two columns of unequal length (which pandas rejects).
    action_rows: list[dict] = []
    recommendations = bundle.narrative.recommendations
    risks = bundle.narrative.risks

    for index, recommendation in enumerate(recommendations):
        action_rows.append(
            {
                "Recommendation": recommendation,
                "Caveat": risks[index] if index < len(risks) else "",
            }
        )
    for extra in risks[len(recommendations) :]:
        action_rows.append({"Recommendation": "", "Caveat": extra})

    actions = pd.DataFrame(action_rows, columns=["Recommendation", "Caveat"])

    data = bundle.df.copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        kpis.to_excel(writer, sheet_name="KPIs", index=False)
        findings.to_excel(writer, sheet_name="Findings", index=False)
        if not actions.empty:
            actions.to_excel(writer, sheet_name="Actions", index=False)
        data.to_excel(writer, sheet_name="Data", index=False)

        book = writer.book
        for name in book.sheetnames:
            _style_sheet(book[name])

        # Summary sheet: wrap the long text cells, highlight the title cell.
        sheet = book["Summary"]
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 110
        sheet["A1"].font = TITLE_FONT
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[0].font = Font(bold=True, color="33333D", name="Segoe UI", size=10)

    return output
